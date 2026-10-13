import io, zipfile
from docx import Document
from ..models import Segment

def extract_docx(filename: str, content: bytes) -> tuple[list[Segment],list[str]]:
    doc=Document(io.BytesIO(content)); seg=[]; idx=0
    for p in doc.paragraphs:
        if p.text.strip():
            idx+=1; seg.append(Segment(filename,f"paragraph:{idx} ({p.style.name})","python-docx",p.text))
    for ti,table in enumerate(doc.tables,1):
        for ri,row in enumerate(table.rows,1):
            text=" | ".join(c.text.replace("\n"," / ") for c in row.cells)
            if text.strip(): seg.append(Segment(filename,f"table:{ti}/row:{ri}","python-docx",text))
    errors=[]
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        images=[n for n in z.namelist() if n.startswith("word/media/")]
    if images: errors.append(f"{len(images)} embedded image(s) detected; image OCR is not implemented for DOCX. Review required.")
    return seg,errors
