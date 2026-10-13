import io
from pptx import Presentation
from ..models import Segment

def extract_pptx(filename: str, content: bytes) -> tuple[list[Segment],list[str]]:
    prs=Presentation(io.BytesIO(content)); seg=[]; errors=[]
    for si,slide in enumerate(prs.slides,1):
        for shi,shape in enumerate(slide.shapes,1):
            loc=f"slide:{si}/shape:{shi}"
            if shape.has_text_frame:
                text="\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if text: seg.append(Segment(filename,loc,"python-pptx",text))
            if shape.has_table:
                for ri,row in enumerate(shape.table.rows,1):
                    seg.append(Segment(filename,f"slide:{si}/table:{shi}/row:{ri}","python-pptx"," | ".join(c.text for c in row.cells)))
            if shape.shape_type==13: errors.append(f"{loc}: embedded image detected; image OCR is not implemented for PPTX.")
    return seg,errors
